from datetime import datetime
import io
import json
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import matplotlib.pyplot as plt
import streamlit as st

aylar = {
    1: "Ocak",
    2: "Şubat",
    3: "Mart",
    4: "Nisan",
    5: "Mayıs",
    6: "Haziran",
    7: "Temmuz",
    8: "Ağustos",
    9: "Eylül",
    10: "Ekim",
    11: "Kasım",
    12: "Aralık",
}
bugun = datetime.now()
bugun_ay_yil = f"{aylar[bugun.month]} {bugun.year}"


def iklim_verisini_yukle():
  dosya_adi = "iklim_verileri.json"
  if os.path.exists(dosya_adi):
    with open(dosya_adi, "r", encoding="utf-8") as f:
      return json.load(f)
  else:
    return {
        "Ankara": {
            "Çankaya": {
                "kis_kt": -12.0,
                "kis_yt": -13.2,
                "yaz_kt": 33.0,
                "yaz_yt": 19.0,
                "enlem": "39° 57' Kuzey",
                "boylam": "32° 53' Doğu",
                "rakim": 949,
                "gsf": 15.5,
            }
        },
        "İstanbul": {
            "Kadıköy": {
                "kis_kt": -2.0,
                "kis_yt": -3.5,
                "yaz_kt": 31.0,
                "yaz_yt": 23.0,
                "enlem": "40° 59' Kuzey",
                "boylam": "29° 02' Doğu",
                "rakim": 30,
                "gsf": 9.0,
            }
        },
    }


def add_toc(paragraph):
  run = paragraph.add_run()
  fldChar1 = OxmlElement("w:fldChar")
  fldChar1.set(qn("w:fldCharType"), "begin")
  instrText = OxmlElement("w:instrText")
  instrText.set(qn("xml:space"), "preserve")
  instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
  fldChar2 = OxmlElement("w:fldChar")
  fldChar2.set(qn("w:fldCharType"), "separate")
  fldChar3 = OxmlElement("w:fldChar")
  fldChar3.set(qn("w:fldCharType"), "end")
  r = run._r
  r.append(fldChar1)
  r.append(instrText)
  r.append(fldChar2)
  r.append(fldChar3)


def _toplu_checkbox_ayarla(anahtarlar, durum):
  for anahtar in anahtarlar:
    st.session_state[anahtar] = durum


def _toplu_secim_butonlari(
    anahtarlar, kolon_basliklari=("☑ TÜMÜNÜ SEÇ", "☐ TÜMÜNÜ KALDIR")
):
  c1, c2 = st.columns(2)
  with c1:
    st.button(
        kolon_basliklari[0],
        key=f"toplu_sec_{anahtarlar[0]}",
        on_click=_toplu_checkbox_ayarla,
        args=(anahtarlar, True),
        use_container_width=True,
    )
  with c2:
    st.button(
        kolon_basliklari[1],
        key=f"toplu_kaldir_{anahtarlar[0]}",
        on_click=_toplu_checkbox_ayarla,
        args=(anahtarlar, False),
        use_container_width=True,
    )


POZ_POMPA_TABLOSU = [
    {
        "poz": "25.360.1301",
        "qmin": 5.0,
        "qmax": 10.0,
        "hmin": 5.0,
        "hmax": 10.0,
        "tanim": (
            "Debisi 5,0-10 m³/h, basıncı 5,0-10 mSS parçalayıcı bıçaklı dalgıç"
            " tip pis su pompası"
        ),
    },
    {
        "poz": "25.360.1302",
        "qmin": 5.0,
        "qmax": 10.0,
        "hmin": 10.0,
        "hmax": 15.0,
        "tanim": (
            "Debisi 5,0-10 m³/h, basıncı 10-15 mSS parçalayıcı bıçaklı dalgıç"
            " tip pis su pompası"
        ),
    },
    {
        "poz": "25.360.1303",
        "qmin": 5.0,
        "qmax": 10.0,
        "hmin": 15.0,
        "hmax": 20.0,
        "tanim": (
            "Debisi 5,0-10 m³/h, basıncı 15-20 mSS parçalayıcı bıçaklı dalgıç"
            " tip pis su pompası"
        ),
    },
    {
        "poz": "25.360.1304",
        "qmin": 10.0,
        "qmax": 15.0,
        "hmin": 5.0,
        "hmax": 10.0,
        "tanim": (
            "Debisi 10-15 m³/h, basıncı 5,0-10 mSS parçalayıcı bıçaklı dalgıç"
            " tip pis su pompası"
        ),
    },
    {
        "poz": "25.360.1305",
        "qmin": 10.0,
        "qmax": 15.0,
        "hmin": 10.0,
        "hmax": 15.0,
        "tanim": (
            "Debisi 10-15 m³/h, basıncı 10-15 mSS parçalayıcı bıçaklı dalgıç"
            " tip pis su pompası"
        ),
    },
    {
        "poz": "25.360.1306",
        "qmin": 15.0,
        "qmax": 20.0,
        "hmin": 5.0,
        "hmax": 10.0,
        "tanim": (
            "Debisi 15-20 m³/h, basıncı 5,0-10 mSS parçalayıcı bıçaklı dalgıç"
            " tip pis su pompası"
        ),
    },
    {
        "poz": "25.360.1307",
        "qmin": 15.0,
        "qmax": 20.0,
        "hmin": 10.0,
        "hmax": 15.0,
        "tanim": (
            "Debisi 15-20 m³/h, basıncı 10-15 mSS parçalayıcı bıçaklı dalgıç"
            " tip pis su pompası"
        ),
    },
    {
        "poz": "25.360.1308",
        "qmin": 15.0,
        "qmax": 20.0,
        "hmin": 15.0,
        "hmax": 20.0,
        "tanim": (
            "Debisi 15-20 m³/h, basıncı 15-20 mSS parçalayıcı bıçaklı dalgıç"
            " tip pis su pompası"
        ),
    },
]


def pompa_pozu_sec(q_m3h, h_mss):
  for kayit in POZ_POMPA_TABLOSU:
    if (kayit["qmin"] <= q_m3h <= kayit["qmax"]) and (
        kayit["hmin"] <= h_mss <= kayit["hmax"]
    ):
      return kayit["poz"], kayit["tanim"], "UYGUN"
  return "SINIR DIŞI", "Çalışma noktası sınır dışındadır!", "GECERSIZ"


def pompa_hidrolik_hesap(q_m3h, h_mss, pompa_verimi=0.60, motor_verimi=0.90):
  rho = 1000.0
  g = 9.81
  p_hid_kw = rho * g * (q_m3h / 3600.0) * h_mss / 1000.0
  p_elektrik_kw = (p_hid_kw / pompa_verimi) / motor_verimi
  motor_kademeleri = [
      0.75,
      1.1,
      1.5,
      2.2,
      3.0,
      4.0,
      5.5,
      7.5,
      11.0,
      15.0,
      18.5,
      22.0,
      30.0,
  ]
  motor_secim = next(
      (x for x in motor_kademeleri if x >= p_elektrik_kw), motor_kademeleri[-1]
  )
  return {
      "q_m3h": q_m3h,
      "h_mss": h_mss,
      "q_lps": q_m3h / 3.6,
      "motor_secim_kw": motor_secim,
  }


URETICI_POMPA_VERITABANI = [
    {
        "marka": "Grundfos",
        "seri": "SEG",
        "model": "SEG.40.15.1",
        "q_min": 0.0,
        "q_max": 5.2,
        "h_max": 26.0,
        "p2_kw": 1.5,
        "curve": [
            (0.0, 26.0),
            (1.0, 24.0),
            (2.0, 21.0),
            (3.0, 17.0),
            (4.0, 13.0),
            (5.0, 9.0),
            (5.2, 8.0),
        ],
        "kaynak": "Grundfos SEG 50 Hz",
    },
    {
        "marka": "Wilo",
        "seri": "Rexa CUT",
        "model": "Rexa CUT GI03.20",
        "q_min": 0.0,
        "q_max": 20.0,
        "h_max": 20.0,
        "p2_kw": 1.1,
        "curve": [
            (0.0, 20.0),
            (4.0, 18.0),
            (8.0, 15.0),
            (12.0, 11.5),
            (16.0, 7.5),
            (20.0, 3.0),
        ],
        "kaynak": "Wilo Rexa CUT",
    },
]


def _egri_degeri(curve, q):
  if not curve or q < curve[0][0] or q > curve[-1][0]:
    return None
  for i in range(len(curve) - 1):
    q0, h0 = curve[i]
    q1, h1 = curve[i + 1]
    if q0 <= q <= q1:
      if q1 == q0:
        return h0
      return h0 + ((q - q0) / (q1 - q0)) * (h1 - h0)
  return curve[-1][1]


def pompa_uretici_sec(q_m3h, h_mss, marka_secimi="Otomatik (Wilo + Grundfos)"):
  adaylar = []
  for pompa in URETICI_POMPA_VERITABANI:
    if marka_secimi == "Wilo" and pompa["marka"] != "Wilo":
      continue
    if marka_secimi == "Grundfos" and pompa["marka"] != "Grundfos":
      continue
    h_egri = _egri_degeri(pompa["curve"], q_m3h)
    if h_egri is None or h_egri < h_mss:
      continue
    adaylar.append((h_egri - h_mss, pompa["p2_kw"], pompa))
  if not adaylar:
    return None
  adaylar.sort(key=lambda x: (x[0], x[1]))
  secilen = dict(adaylar[0][2])
  secilen["h_calisma"] = _egri_degeri(secilen["curve"], q_m3h)
  return secilen


def pompa_secim_egrisi(q_m3h, h_mss, marka_secimi="Otomatik (Wilo + Grundfos)"):
  poz, _, durum = pompa_pozu_sec(q_m3h, h_mss)
  if durum == "UYGUN":
    model = pompa_uretici_sec(q_m3h, h_mss, marka_secimi)
    if model:
      return (
          [p[0] for p in model["curve"]],
          [p[1] for p in model["curve"]],
          f"{model['marka']} {model['model']} - Üretici Q-H Eğrisi",
          model,
      )
  return [2, 22], [22, 2], "Sınır Dışı Çalışma Noktası!", None


def pompa_grafigi_png(
    q_egrisi, h_egrisi, q_calisma, h_calisma, baslik, anonim=False
):
  fig, ax = plt.subplots(figsize=(7.0, 3.8))
  ax.plot(
      q_egrisi,
      h_egrisi,
      linewidth=2.0,
      label=("Pompa Performans Eğrisi" if anonim else baslik),
  )
  ax.scatter(
      [q_calisma],
      [h_calisma],
      s=65,
      zorder=5,
      label=f"Çalışma Noktası ({q_calisma:.2f} m³/h, {h_calisma:.2f} mSS)",
  )
  ax.set_xlabel("Debi Q [m³/h]")
  ax.set_ylabel("Basma Yüksekliği H [mSS]")
  ax.set_title("Pompa Performans Eğrisi" if anonim else baslik, fontsize=11)
  ax.grid(True, alpha=0.3)
  ax.legend(fontsize=8)
  fig.tight_layout()
  buf = io.BytesIO()
  fig.savefig(buf, format="png", dpi=180, bbox_inches="tight")
  plt.close(fig)
  buf.seek(0)
  return buf


def rapor_word_stillerini_uygula(doc):
  for style_name in [
      "Normal",
      "Body Text",
      "List Paragraph",
      "List Bullet",
      "List Number",
  ]:
    try:
      stl = doc.styles[style_name]
      stl.font.name = "Times New Roman"
      stl._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
      stl.font.size = Pt(12)
    except KeyError:
      pass
  h1 = doc.styles["Heading 1"]
  h1.font.name = "Times New Roman"
  h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
  h1.font.size = Pt(14)
  h1.font.bold = True
  h1.paragraph_format.page_break_before = True
  for level in [2, 3, 4]:
    h = doc.styles[f"Heading {level}"]
    h.font.name = "Times New Roman"
    h._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h.font.size = Pt(12)
    h.font.bold = True
  for paragraph in doc.paragraphs:
    for run in paragraph.runs:
      run.font.name = "Times New Roman"
      run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
      if paragraph.style and paragraph.style.name.startswith("Heading 1"):
        run.font.size = Pt(14)
        run.font.bold = True
      elif paragraph.style and paragraph.style.name.startswith("Heading"):
        run.font.size = Pt(12)
        run.font.bold = True
      elif run.font.size is None:
        run.font.size = Pt(12)
  for table in doc.tables:
    for row in table.rows:
      for cell in row.cells:
        for paragraph in cell.paragraphs:
          for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(12)
import streamlit as strlit
from ana_modul import _toplu_secim_butonlari, iklim_verisini_yukle

iklim_veritabani = iklim_verisini_yukle()


def bolum_1_5_arayuz():
  strlit.header("1. Kapak Bilgileri")
  strlit.text_input(
      "Şirket / Kuruluş İsmi",
      (
          "FUGA MEKANİK MÜHENDİSLİK MÜŞAVİRLİK İNŞ.SAN.TİC.LTD.ŞTİ"
          if "sirket_adi" not in strlit.session_state
          else strlit.session_state["sirket_adi"]
      ),
      key="sirket_adi",
  )
  strlit.text_input("İşin Adı / Proje Başlığı", key="is_adi")
  strlit.text_input(
      "Rapor Türü",
      "MEKANİK TESİSAT UYGULAMA PROJESİ HESAP RAPORU",
      key="rapor_turu",
  )
  strlit.text_input("Hazırlayan Mühendis", "Mehmet Küçük", key="hazirlayan")
  strlit.text_input("MMO Oda No", "109913", key="mmo_no")
  strlit.text_input("Rapor Tarihi", key="tarih")

  strlit.header("2. UYGULANACAK STANDART VE YÖNETMELİKLER")
  std_keys = [
      "std_ts_825",
      "std_yangin",
      "std_bep_2008_2010",
      "std_ts_1258",
      "std_ts_826",
      "std_ts_2164",
      "std_ts_3419",
      "std_ts_en_12056_2",
      "std_ts_en_12845",
      "std_mmo_84",
      "std_mmo_352_5",
      "std_mmo_122",
      "std_mmo_133",
      "std_mmo_155",
      "std_ashrae",
      "std_su",
      "std_klima",
      "std_tesisat",
      "std_kanal",
      "std_asansor",
      "std_deprem",
      "std_akustik",
      "std_isg",
  ]
  _toplu_secim_butonlari(std_keys)

  strlit.checkbox(
      "TS 825 - BİNALARDA ISI YALITIM KURALLARI",
      key="std_ts_825",
      value=True,
  )
  strlit.checkbox(
      '09 Eylül 2009 tarih ve 27344 numaralı sayısında yayımlanan " BİNALARIN'
      ' YANGINDAN KORUNMASI HAKKINDA YÖNETMELİK"',
      key="std_yangin",
      value=True,
  )
  strlit.checkbox(
      "5 Aralık 2008 tarih, 27075 sayılı resmi gazetede yayımlanan “BİNALARDA"
      " ENERJİ PERFORMANSI YÖNETMELİĞİ” ve 1 Nisan 2010 tarih, 27539 sayılı resmi"
      " gazetede yayımlanan “BİNALARDA ENERJİ PERFORMANSI YÖNETMELİĞİ”",
      key="std_bep_2008_2010",
      value=True,
  )
  strlit.checkbox(
      "TS 1258 – TEMİZSU TESİsATI HESAP KURALLARI",
      key="std_ts_1258",
      value=True,
  )
  strlit.checkbox(
      "TS 826 – BİNALARDA PİSSU TESİsATI HESAPLAMA KURALLARI",
      key="std_ts_826",
      value=True,
  )
  strlit.checkbox(
      "TS 2164 - KALORİFER TESİsATI PROJELENDİRME KURALLARI",
      key="std_ts_2164",
      value=True,
  )
  strlit.checkbox(
      "TS 3419 – HAVALANDIRMA VE İKLİMLENDİRME TESİSLERİ PROJELENDİRME"
      " KURALLARI",
      key="std_ts_3419",
      value=True,
  )
  strlit.checkbox(
      "TS EN 12056-2 – CAZİBELİ DRENAJ SİSTEMLERİ -BİNA İÇİ- TASARIM VE"
      " HESAPLAMA",
      key="std_ts_en_12056_2",
      value=True,
  )
  strlit.checkbox(
      "TS EN 12845 – SABİT YANGIN SÖNDÜRME SİSTEMLERİ – OTOMATİK SPRİNKLER"
      " SİSTEMLERİ- TASARIM, MONTAJ VE BAKIM",
      key="std_ts_en_12845",
      value=True,
  )
  strlit.checkbox(
      "MMO KALORİFER TESİsATI PROJE HAZIRLAMA ESASLARI(Y.NO:84)",
      key="std_mmo_84",
      value=True,
  )
  strlit.checkbox(
      "MMO KALORİFER TESİsATI (Y.NO:352/5)", key="std_mmo_352_5", value=True
  )
  strlit.checkbox(
      "MMO SIHHİ TESİSAT PROJE HAZIRLAMA ESASLARI(Y.NO:122)",
      key="std_mmo_122",
      value=True,
  )
  strlit.checkbox(
      "MMO GAZ TESİsATI PROJE HAZIRLAMA ESASLARI(Y.NO:133)",
      key="std_mmo_133",
      value=True,
  )
  strlit.checkbox(
      "MMO KAZAN VE BACA(Y.NO:155)", key="std_mmo_155", value=True
  )
  strlit.checkbox(
      "ASHRAE Standartları", key="std_ashrae", value=True
  )
  strlit.checkbox(
      "İçmesuyu Temizleme ve Dağıtım Sistemleri Standartları",
      key="std_su",
      value=True,
  )
  strlit.checkbox(
      "Klima ve Havalandırma Tesisatı Yönetmelikleri",
      key="std_klima",
      value=True,
  )
  strlit.checkbox(
      "Merkezi Isıtma ve Sıhhi Sıcak Su Sistemlerinde Isı Maliyetlerinin"
      " Paylaştırılmasına İlişkin Yönetmelik",
      key="std_tesisat",
      value=False,
  )
  strlit.checkbox(
      "Kanalizasyon Şebekesi Olmayan Yerlerde Yapılacak Çukurlar",
      key="std_kanal",
      value=False,
  )
  strlit.checkbox(
      "Asansör Yönetmeliği ve İlgili Standartlar",
      key="std_asansor",
      value=False,
  )
  strlit.checkbox(
      "Türkiye Bina Deprem Yönetmeliği (Mekanik Ekipman Askı ve Destekleri)",
      key="std_deprem",
      value=True,
  )
  strlit.checkbox(
      "Binaların Gürültüye Karşı Korunması Yönetmeliği",
      key="std_akustik",
      value=False,
  )
  strlit.checkbox(
      "İş Sağlığı ve Güvenliği Kanunu ve İlgili Yönetmelikler",
      key="std_isg",
      value=True,
  )
  strlit.text_area("Eklemek istediğiniz ilave standartlar", key="ek_standartlar")

  strlit.header("3. MEKANİK TESİSAT PROJE KAPSAMI")
  kapsam_keys = [
      "kapsam_isitma",
      "kapsam_sogutma",
      "kapsam_soguk_su",
      "kapsam_sicak_su",
      "kapsam_yangin_depo",
      "kapsam_atik_su",
      "kapsam_yangin_dagitim",
      "kapsam_kazan_dairesi",
      "kapsam_havalandirma",
      "kapsam_basinc_hava",
      "kapsam_medikal_gaz",
      "kapsam_otomatik",
  ]
  _toplu_secim_butonlari(kapsam_keys)
  strlit.checkbox("Isıtma tesisatı,", key="kapsam_isitma", value=True)
  strlit.checkbox("Soğutma tesisatı,", key="kapsam_sogutma", value=True)
  strlit.checkbox(
      "Kullanma soğuk suyu tesisatı,", key="kapsam_soguk_su", value=True
  )
  strlit.checkbox(
      "Kullanma sıcak suyu tesisatı,", key="kapsam_sicak_su", value=True
  )
  strlit.checkbox(
      "Yangın ve kullanma suyu depolaması ve dağıtımı,",
      key="kapsam_yangin_depo",
      value=True,
  )
  strlit.checkbox(
      "Yapı içinde atık su tesisatı (Yapı çıkış rögarına),",
      key="kapsam_atik_su",
      value=True,
  )
  strlit.checkbox(
      "Yangın suyu iç ve dış dağıtım sistemleri,",
      key="kapsam_yangin_dagitim",
      value=True,
  )
  strlit.checkbox(
      "Merkezi ısıtma kazan dairesi ve tali teknik hacimler,",
      key="kapsam_kazan_dairesi",
      value=True,
  )
  strlit.checkbox("Havalandırma Tesisatı", key="kapsam_havalandirma", value=True)
  strlit.checkbox(
      "Basınçlı hava tesisatı,", key="kapsam_basinc_hava", value=False
  )
  strlit.checkbox("Medikal gaz tesisatı", key="kapsam_medikal_gaz", value=False)
  strlit.checkbox(
      "Otomatik kontrol sistemi kavramı tanımı,",
      key="kapsam_otomatik",
      value=True,
  )
  strlit.text_area("İlave proje kapsam maddeleri", key="ek_kapsam")

  strlit.header("5. İKLİM, KONFOR ŞARTLARI VE TASARIM KRİTERLERİ")
  iller = sorted(list(iklim_veritabani.keys()))
  secilen_il = strlit.selectbox("İl seçin:", iller, key="secilen_il")
  ilceler = sorted(list(iklim_veritabani[secilen_il].keys()))
  secilen_ilce = strlit.selectbox("İlçe seçin:", ilceler, key="secilen_ilce")
  return iklim_veritabani[secilen_il][secilen_ilce]
import math
import streamlit as strlit
from ana_modul import _toplu_secim_butonlari, pompa_grafigi_png, pompa_hidrolik_hesap, pompa_pozu_sec, pompa_secim_egrisi


def bolum_6_arayuz():
  strlit.header("6. SIHHİ TESİSAT")
  strlit.subheader("6.1 SIHHİ TESİSAT ÖN BİLGİLER")
  sihhi_keys = [
      "sih_sec_1",
      "sih_sec_2",
      "sih_sec_3",
      "sih_sec_4",
      "sih_sec_depo_tipi",
      "sih_sec_5",
      "sih_sec_6",
      "sih_sec_7",
      "sih_sec_8",
      "sih_sec_9",
      "sih_sec_10",
      "sih_sec_12",
  ]
  _toplu_secim_butonlari(sihhi_keys)

  strlit.checkbox(
      "Bütün tesisin kullanma soğuk su ihtiyacı şehir şebekesinden"
      " sağlanacaktır.",
      key="sih_sec_1",
      value=True,
  )
  strlit.checkbox(
      "Bütün tesisin kullanma soğuk su ihtiyacı kampüs içi su deposu dağıtım"
      " hattından sağlanacaktır.",
      key="sih_sec_2",
      value=False,
  )
  strlit.checkbox(
      "Temiz su boru çapları yükleme birimine verilmiştir.",
      key="sih_sec_3",
      value=True,
  )
  strlit.checkbox(
      "Bütün binanın kullanma soğuk su ihtiyacı soğuk su deposundan"
      " sağlanacaktır.",
      key="sih_sec_4",
      value=True,
  )
  strlit.multiselect(
      "Soğuk Su Deposu Konumu:",
      ["Bodrum kat", "Zemin kat", "1. kat", "2. kat", "Çatı katı"],
      default=["Bodrum kat"],
      key="sih_depo_konumlari",
  )
  strlit.checkbox(
      "Bina da kullanım soğuk su depolaması için belirtilen tipte su deposu"
      " kullanılmıştır.",
      key="sih_sec_depo_tipi",
      value=True,
  )
  strlit.multiselect(
      "Kullanma Soğuk Su Deposu Tipi:",
      [
          "Paslanmaz Çelik Modüler su deposu",
          "Galvaniz Çelik Modüler su deposu",
          "GRP (Cam Takviyeli Polyester) Modüler su deposu",
          "Betonarme Su deposu",
          "Silindirik Plastik Su deposu",
      ],
      default=["Paslanmaz Çelik Modüler su deposu"],
      key="sih_depo_tipleri",
  )

  strlit.subheader("6.2 PİS SU TESİSATI ESASLARI")
  pissu_keys = [
      "pissu_sec_1",
      "pissu_sec_2",
      "pissu_sec_3",
      "pissu_sec_4",
      "pissu_sec_5",
      "pissu_sec_6",
      "pissu_sec_7",
  ]
  _toplu_secim_butonlari(pissu_keys)
  strlit.checkbox(
      "Yapının atık suları binanın pik kolonlarla toplanacaktır.",
      key="pissu_sec_1",
      value=True,
  )
  strlit.multiselect(
      "Atık su toplama konumu:",
      ["Bodrum kat", "Zemin kat", "1. kat", "Çatı katı"],
      default=["Bodrum kat"],
      key="pis_su_konumlari",
  )
  strlit.multiselect(
      "Atık su boru geçiş yeri:",
      ["döşemesinden", "tavanından", "asma tavan arasından"],
      default=["döşemesinden"],
      key="pis_su_gecisler",
  )
  strlit.checkbox(
      "Pis su kolonları üzerinde temizleme kapakları vardır.",
      key="pissu_sec_2",
      value=True,
  )
  strlit.checkbox(
      "Teknik hacimlerde ızgaralı kanallar yapılacaktır.",
      key="pissu_sec_3",
      value=True,
  )
  strlit.checkbox(
      "Sessiz PVC boru kullanılacaktır.", key="pissu_sec_4", value=True
  )
  strlit.checkbox(
      "Yükleme birimi yöntemine göre belirlenmiştir.",
      key="pissu_sec_5",
      value=True,
  )
  strlit.checkbox(
      "Kot kurtarmayan katlar bodrum katta pis su çukurunda toplanacaktır.",
      key="pissu_sec_6",
      value=True,
  )
  strlit.checkbox(
      "Rögarlar vasıtasıyla yoldan geçen kanalına bağlanacaktır.",
      key="pissu_sec_7",
      value=True,
  )

  strlit.subheader("6.2.2 Pis Su Terfi Pompaları Hesap Modülü")
  secilen_psp_listesi = strlit.multiselect(
      "Projede yer alacak Pis Su Terfi Pompalarını seçin:",
      [f"PSP-{i:02d}" for i in range(1, 11)],
      default=["PSP-01"],
      key="secilen_psp_listesi",
  )
  strlit.checkbox(
      "Cihaz Poz Numarasını Rapora Aktar", value=True, key="poz_aktar_chk"
  )
  pompa_marka_secimi = strlit.selectbox(
      "Pompa üreticisi / seçim modu",
      ["Otomatik (Wilo + Grundfos)", "Wilo", "Grundfos"],
      key="pompa_marka_secimi",
  )

  psp_parametreleri = {}
  if secilen_psp_listesi:
    for psp in secilen_psp_listesi:
      with strlit.expander(f"⚙️ {psp} Hesap Modülü", expanded=True):
        bina_tipi = strlit.selectbox(
            f"{psp} Bina Kullanım Türü",
            [
                "Evler, oteller, ofisler (Düzensiz) [k=0.5]",
                "Hastaneler, geniş gıda [k=0.7]",
                "Okullar, umumi tuvaletler [k=1.0]",
                "Endüstriyel laboratuvarlar [k=1.2]",
            ],
            key=f"{psp}_bina",
        )
        k_katsayisi = strlit.number_input(
            f"{psp} Eşzamanlık (k)", value=0.5, step=0.05, key=f"{psp}_k"
        )

        ac1, ac2 = strlit.columns(2)
        with ac1:
          adet_hela = strlit.number_input(
              "Hela / Klozet", min_value=0, value=4, key=f"{psp}_hela"
          )
          adet_lavabo = strlit.number_input(
              "Lavabo / Bide", min_value=0, value=6, key=f"{psp}_lav"
          )
        with ac2:
          adet_banyo = strlit.number_input(
              "Küvet / Duş", min_value=0, value=2, key=f"{psp}_ban"
          )
          adet_evye = strlit.number_input(
              "Eviye", min_value=0, value=1, key=f"{psp}_evy"
          )

        yb_toplam = (
            adet_hela * 8
            + adet_lavabo * 2
            + adet_banyo * 7
            + adet_evye * 4
        )
        net_q_m3h = round(k_katsayisi * math.sqrt(max(0, yb_toplam)) * 3.6, 2)

        emniyet_katsayisi = strlit.selectbox(
            "Debi Emniyet Oranı", [1.0, 1.10, 1.15, 1.20, 1.25, 1.30], key=f"{psp}_emn"
        )
        toplam_v_val = strlit.number_input(
            "Toplam Debi [m³/h]",
            value=max(1.0, net_q_m3h * emniyet_katsayisi),
            key=f"{psp}_v_num",
        )
        h_val = strlit.number_input(
            "Basma Yüksekliği [mSS]", value=12.0, key=f"{psp}_h_num"
        )
        asil_adedi = strlit.selectbox(
            "Asıl Pompa Adedi", [1, 2, 3], key=f"{psp}_asil"
        )
        yedek_adedi = strlit.selectbox(
            "Yedek Pompa Adedi", [1, 2], key=f"{psp}_yedek"
        )

        pompa_basina_v = toplam_v_val / asil_adedi
        hesap = pompa_hidrolik_hesap(pompa_basina_v, h_val)
        poz, tanim, durum = pompa_pozu_sec(pompa_basina_v, h_val)
        q_egr, h_egr, baslik, secilen_uretici = pompa_secim_egrisi(
            pompa_basina_v, h_val, pompa_marka_secimi
        )

        if durum == "UYGUN":
          strlit.success(f"✅ Uygun Poz: {poz}")
        else:
          strlit.error("❌ Çalışma noktası sınır dışı!")

        strlit.image(
            pompa_grafigi_png(q_egr, h_egr, pompa_basina_v, h_val, baslik),
            use_container_width=True,
        )

        psp_parametreleri[psp] = {
            "bina_tipi": bina_tipi,
            "k_katsayisi": k_katsayisi,
            "toplam_yb": yb_toplam,
            "net_q_m3h": net_q_m3h,
            "v_toplam": toplam_v_val,
            "v_tek": pompa_basina_v,
            "h": h_val,
            "guc": hesap["motor_secim_kw"],
            "asil_adedi": asil_adedi,
            "adet_str": f"{asil_adedi + yedek_adedi} ({asil_adedi} Asıl + {yedek_adedi} Yedek)",
            "poz": poz,
            "poz_durumu": durum,
            "poz_tanim": tanim,
            "pompa_curve": list(zip(q_egr, h_egr)),
            "pompa_egrisi_basligi": baslik,
            "tablo_satirlari": [
                ("Hela / Klozet", 8, adet_hela, adet_hela * 8),
                ("Lavabo / Bide", 2, adet_lavabo, adet_lavabo * 2),
                ("Küvet / Duş", 7, adet_banyo, adet_banyo * 7),
                ("Eviye", 4, adet_evye, adet_evye * 4),
            ],
        }

  strlit.subheader("6.3 SIHHİ TESİSAT CİHAZ SEÇİMLERİ")
  strlit.subheader("6.3.1 KULLANMA SOĞUK SUYU DEPOSU SEÇİMİ")
  strlit.checkbox(
      "Su deposu kapasite hesaplamalarında kişi başı günlük su tüketim"
      " miktarı, kişi sayısı ve binanın kullanım amacı göz önüne alınmıştır.",
      key="depo_sec_1",
      value=True,
  )
  strlit.checkbox(
      "Su deposu içerisinde su kalitesinin korunması ve ölü hacim oluşumunun"
      " önlenmesi için bölme perdeleri yer alacaktır.",
      key="depo_sec_2",
      value=True,
  )
  strlit.checkbox(
      "Su deposunda taşma, deşarj, havalandırma boruları ile bakım ve temizlik"
      " için adam geçiş kapağı (manhole) bulunacaktır.",
      key="depo_sec_3",
      value=True,
  )
  strlit.text_area("İlave Depo Notu", key="ek_depo_notu")

  return psp_parametreleri
import io
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
import streamlit as strlit

from ana_modul import add_toc, bugun_ay_yil, pompa_grafigi_png, rapor_word_stillerini_uygula
from bolum_1_5 import bolum_1_5_arayuz
from bolum_6 import bolum_6_arayuz

strlit.set_page_config(
    page_title="Mühendislik Proje Raporu Otomasyonu", layout="wide"
)

# --- PROJE YÖNETİMİ & FARKLI KAYDET (SIDEBAR) ---
strlit.sidebar.header("📁 Proje Yönetimi")

if "proje_verileri" not in strlit.session_state:
  strlit.session_state["proje_verileri"] = {}

proje_no = strlit.sidebar.text_input("Proje Numarası / Kodu", "PRJ-2026-001")
proje_adi_input = strlit.sidebar.text_input("Proje Adı", "Örnek Bina Projesi")


def proje_verilerini_topla():
  veri = {}
  for k, v in strlit.session_state.items():
    if isinstance(v, (str, int, float, bool, list)):
      veri[k] = v
  return veri


proje_json = json.dumps(proje_verilerini_topla(), ensure_ascii=False, indent=4)
strlit.sidebar.download_button(
    label="💾 Projeyi Kaydet (JSON)",
    data=proje_json,
    file_name=f"{proje_no}_{proje_adi_input}.json".replace(" ", "_"),
    mime="application/json",
)

yuklenen_dosya = strlit.sidebar.file_uploader(
    "📂 Kayıtlı Proje Yükle (.json)", type=["json"]
)
if yuklenen_dosya is not None:
  try:
    yuklenen_veri = json.load(yuklenen_dosya)
    for k, v in yuklenen_veri.items():
      strlit.session_state[k] = v
    strlit.sidebar.success("Proje başarıyla yüklendi!")
  except Exception as e:
    strlit.sidebar.error(f"Dosya okunamadı: {e}")

strlit.title("Mühendislik Proje Raporu Otomasyonu")
strlit.write(
    "Modüler Mimari & Proje Yönetim Paneli ile çalışmaktasınız. Sol panelden"
    " projelerinizi kaydedebilir veya başka bir projeyi yükleyebilirsiniz."
)

iklim_veri = bolum_1_5_arayuz()
psp_parametreleri = bolum_6_arayuz()

if strlit.button("Raporu Oluştur (.docx)"):
  gecersiz_var = any(
      p["poz_durumu"] != "UYGUN" for p in psp_parametreleri.values()
  )
  if gecersiz_var:
    strlit.error(
        "❌ Rapor oluşturulamadı! Seçilen pompalardan biri sınır dışındadır."
    )
  else:
    aktif_is = strlit.session_state.get("is_adi", "")
    aktif_sirket = strlit.session_state.get(
        "sirket_adi", "FUGA MEKANİK MÜHENDİSLİK..."
    )

    doc = Document()
    cover_section = doc.sections[0]
    cover_section.top_margin = Inches(1.5)
    cover_section.bottom_margin = Inches(1.5)
    cover_section.left_margin = Inches(1.2)
    cover_section.right_margin = Inches(1.2)

    p_sirket = doc.add_paragraph()
    p_sirket.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_s = p_sirket.add_run(aktif_sirket.upper())
    run_s.font.size = Pt(13)
    run_s.font.bold = True

    doc.add_paragraph()
    doc.add_paragraph()
    if aktif_is:
      p_is = doc.add_paragraph()
      p_is.alignment = WD_ALIGN_PARAGRAPH.CENTER
      p_is.add_run("PROJE ADI:\n").font.size = Pt(11)
      run_i = p_is.add_run(aktif_is)
      run_i.font.size = Pt(16)
      run_i.font.bold = True
      doc.add_paragraph()

    p_tur = doc.add_paragraph()
    p_tur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t = p_tur.add_run(
        strlit.session_state.get(
            "rapor_turu", "MEKANİK TESİSAT UYGULAMA PROJESİ"
        ).upper()
    )
    run_t.font.size = Pt(14)
    run_t.font.bold = True

    for _ in range(4):
      doc.add_paragraph()
    p_alt = doc.add_paragraph()
    p_alt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_alt.add_run(
        f"Hazırlayan:\n{strlit.session_state.get('hazirlayan', 'Mehmet Küçük')} (Makine"
        f" Mühendisi)\nMMO Oda No:"
        f" {strlit.session_state.get('mmo_no', '109913')}\n\nTarih:\n{strlit.session_state.get('tarih', bugun_ay_yil)}"
    )

    doc.add_page_break()
    doc.add_heading("İÇİNDEKİLER", level=1)
    add_toc(doc.add_paragraph())

    doc.add_page_break()
    body_sec = doc.add_section()
    body_sec.top_margin = Inches(1.2)
    body_sec.bottom_margin = Inches(1.2)
    body_sec.left_margin = Inches(1.2)
    body_sec.right_margin = Inches(1.2)

    doc.add_heading("1. GENEL BİLGİLER", level=1)
    doc.add_paragraph(
        f"Bu raporda '{aktif_is}' için tasarlanan mekanik tesisatlar"
        " açıklanmıştır."
    )
    doc.add_paragraph(
        f"Yapı {strlit.session_state.get('secilen_il', '')} ili"
        f" {strlit.session_state.get('secilen_ilce', '')} ilçesinde inşa"
        " edilecektir."
    )

    doc.add_heading("5. İKLİM, KONFOR ŞARTLARI VE TASARIM KRİTERLERİ", level=1)
    doc.add_heading("5.1 DIŞ HAVA TASARIM KRİTERLERİ", level=2)
    doc.add_paragraph(
        f"• KIŞ: {iklim_veri['kis_kt']} °C KT , {iklim_veri['kis_yt']} °C YT"
    )
    doc.add_paragraph(
        f"• YAZ: {iklim_veri['yaz_kt']} °C KT , {iklim_veri['yaz_yt']} °C YT"
    )
    doc.add_paragraph(f"• Enlem: {iklim_veri['enlem']}")
    doc.add_paragraph(f"• Boylam: {iklim_veri['boylam']}")
    doc.add_paragraph(f"• Rakım: {iklim_veri['rakim']} m.")

    doc.add_heading("6. SIHHİ TESİSAT", level=1)
    doc.add_heading("6.3 SIHHİ TESİSAT CİHAZ SEÇİMLERİ", level=1)
    doc.add_heading("6.3.1 KULLANMA SOĞUK SUYU DEPOSU SEÇİMİ", level=2)

    depo_tipleri = strlit.session_state.get("sih_depo_tipleri", [])
    dinamik_tip = (
        depo_tipleri[0].lower() if depo_tipleri else "modüler su deposu"
    )

    doc.add_paragraph(
        "Binanın kullanma soğuk suyu ihtiyacının karşılanması ve kesintilere"
        f" karşı güvence altına alınması amacıyla {dinamik_tip} tasarlanmıştır."
    )
    if strlit.session_state.get("depo_sec_1", True):
      doc.add_paragraph(
          "Su deposu kapasite hesaplamalarında kişi başı günlük su tüketim"
          " miktarı, kişi sayısı ve binanın kullanım amacı faktörleri göz önüne"
          " alınmıştır."
      )
    if strlit.session_state.get("depo_sec_2", True):
      doc.add_paragraph(
          "Su deposu içerisinde su kalitesinin korunması ve ölü hacim oluşumunun"
          " önlenmesi için bölme perdeleri yer alacaktır."
      )
    if strlit.session_state.get("depo_sec_3", True):
      doc.add_paragraph(
          "Su deposunda taşma, deşarj, havalandırma boruları ile bakım ve temizlik"
          " için adam geçiş kapağı (manhole) bulunacaktır."
      )

    rapor_word_stillerini_uygula(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    strlit.success("Modüler rapor başarıyla hazırlandı!")
    strlit.download_button(
        label="📥 Word Dosyasını İndir (.docx)",
        data=buffer,
        file_name=f"{proje_no}_{aktif_is.replace(' ', '_')}_Rapor.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )
